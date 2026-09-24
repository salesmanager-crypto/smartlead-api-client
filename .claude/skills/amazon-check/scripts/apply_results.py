#!/usr/bin/env python3
"""Write Amazon check results into the lead sheet, one result per brand applied to every contact row.

Reads queue.json (from build_queue.py, for the row mapping) and out.json (from verify_brands.py).
Adds columns after the last existing column unless --after is given. Never reorders rows.

Columns written: Amazon Status | Sold By | Amazon Brand Name | Name Match? | Amazon URL | Brand Store URL | Amazon Notes
With --rename-company: also adds "Original Company" after Company and sets Company = Amazon Brand Name
on rows whose status is Sells or Reseller / Unclear and whose brand was found (never on Does Not Sell).

Usage: apply_results.py leads.xlsx --queue queue.json --results out.json --out "Leads - Amazon Checked.xlsx" [--sheet NAME] [--after "Amazon Notes"] [--rename-company]
"""
import argparse, json, re, copy, collections, openpyxl
from openpyxl.utils import get_column_letter as L

p = argparse.ArgumentParser()
p.add_argument("xlsx"); p.add_argument("--queue", required=True); p.add_argument("--results", required=True)
p.add_argument("--out", required=True); p.add_argument("--sheet"); p.add_argument("--after"); p.add_argument("--rename-company", action="store_true")
a = p.parse_args()

def norm(s): return re.sub(r'[^a-z0-9]', '', (s or '').lower())
SUF = r'\b(inc|llc|co|corp|corporation|ltd|limited|company|the|group|brands?|international|usa|designs?|apparel|clothing)\b'
def core(s): return norm(re.sub(SUF, '', (s or '').lower()))

queue = json.load(open(a.queue)); res = json.load(open(a.results))
wb = openpyxl.load_workbook(a.xlsx); ws = wb[a.sheet] if a.sheet else wb[wb.sheetnames[0]]
H = [c.value for c in ws[1]]
hdr_style = copy.copy(ws.cell(1, 1)._style)

def insert_after(header, newname):
    global H
    i = (H.index(header) + 1 if header else len(H)) + 1
    ws.insert_cols(i); c = ws.cell(1, i, newname); c._style = copy.copy(hdr_style); H = [c.value for c in ws[1]]
    return i

last = max(r for r in range(2, ws.max_row + 1) if any(c.value not in (None, '') for c in ws[r]))
cC = H.index('Company') + 1 if 'Company' in H else None
if a.rename_company and cC:
    cOC = insert_after('Company', 'Original Company'); cC = H.index('Company') + 1
    for r in range(2, last + 1): ws.cell(r, cOC).value = ws.cell(r, cC).value; ws.cell(r, cOC)._style = copy.copy(ws.cell(r, cC)._style)

new = ['Amazon Status', 'Sold By', 'Amazon Brand Name', 'Name Match?', 'Amazon URL', 'Brand Store URL', 'Amazon Notes']
anchor = a.after
cols = {}
for n in new:
    cols[n] = insert_after(anchor, n); anchor = n
for n, w in zip(new, (18, 22, 26, 12, 45, 45, 55)): ws.column_dimensions[L(cols[n])].width = w

counts = collections.Counter(); changed = collections.Counter(); notfound = []
for b in queue:
    o = res.get(b['key'], {})
    brand = o.get('brand') or ''
    status = o.get('status') or ('Sells' if o.get('sold') in ('Brand (Seller Central)', 'Amazon (Vendor)') and (o.get('store_url') or o.get('evidence')) else 'Reseller / Unclear' if brand else 'Does Not Sell')
    url = o.get('store_url') or o.get('evidence') or b.get('link') or ''
    comp = b['company'] or b['domain'] or b['key']
    match = 'Yes' if brand and core(brand) == core(comp) else ('No' if brand else '')
    for r in b['rows']:
        ref = copy.copy(ws.cell(r, cC)._style) if cC else None
        vals = {'Amazon Status': status, 'Sold By': o.get('sold') or ('Unknown' if brand else ''), 'Amazon Brand Name': brand or ('NOT FOUND' if status != 'Does Not Sell' else ''),
                'Name Match?': match, 'Amazon URL': url or None, 'Brand Store URL': o.get('store_url') or None, 'Amazon Notes': o.get('notes') or ''}
        for n, v in vals.items():
            c = ws.cell(r, cols[n]); c.value = v
            if ref: c._style = copy.copy(ref)
            if n in ('Amazon URL', 'Brand Store URL') and v: c.hyperlink = v; c.style = 'Hyperlink'
        if a.rename_company and cC and brand and status != 'Does Not Sell':
            old = ws.cell(r, cC).value
            if (old or '') != brand: changed[(old, brand)] += 1
            ws.cell(r, cC).value = brand
        counts[status] += 1
    if not brand and status != 'Does Not Sell': notfound.append(comp)

wb.save(a.out)
print(f"saved {a.out}; data rows {last - 1}; status counts {dict(counts)}")
if a.rename_company:
    print("\nCOMPANY CHANGES (Original -> Amazon):"); [print(f"  {o} -> {n}  [{k}]") for (o, n), k in sorted(changed.items(), key=lambda x: str(x[0][0]).lower())]
print("\nNOT FOUND:"); [print("  ", x) for x in sorted(notfound, key=str.lower)]
